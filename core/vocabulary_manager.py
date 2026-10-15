import os
import json
import csv
from openpyxl import Workbook
from datetime import datetime

class VocabularyManager:
    def __init__(self, data_file="data/vocabulary.json"):
        self.data_file = data_file
        os.makedirs(os.path.dirname(data_file), exist_ok=True)
        self.words = self._load_words()

    def _load_words(self):
        if os.path.exists(self.data_file):
            with open(self.data_file, 'r', encoding='utf-8') as f:
                return json.load(f)
        return {}

    def _save(self):
        with open(self.data_file, 'w', encoding='utf-8') as f:
            json.dump(self.words, f, ensure_ascii=False, indent=2)

    def add_word(self, word, base_form="", definition="", source_article="", detail=None):
        key = base_form if base_form else word.lower()
        if key not in self.words:
            self.words[key] = {
                "base_form": key,
                "definition": definition,
                "source_articles": [source_article] if source_article else [],
                "add_time": datetime.now().isoformat(),
                "mastery_level": 0,
                "detail": detail or {},
                "occurrences": []
            }
        else:
            if source_article and source_article not in self.words[key]["source_articles"]:
                self.words[key]["source_articles"].append(source_article)
        self._save()
        return key

    def get_word(self, word):
        key = word.lower()
        return self.words.get(key, None)

    def get_article_words(self, article_id):
        result = []
        for word, info in self.words.items():
            if article_id in info["source_articles"]:
                result.append(info)
        return result

    def get_all_words(self):
        return list(self.words.values())

    def update_mastery(self, word, level):
        key = word.lower()
        if key in self.words:
            self.words[key]["mastery_level"] = level
            self._save()

    def delete_word(self, word):
        key = word.lower()
        if key in self.words:
            del self.words[key]
            self._save()

    def import_from_file(self, file_path, group_name="", local_dict=None):
        ext = os.path.splitext(file_path)[1].lower()
        imported = 0
        
        def get_def(word, original_def=""):
            if original_def:
                return original_def
            if local_dict:
                return local_dict.lookup(word.lower()) or ""
            return ""
        
        if ext == '.txt':
            with open(file_path, 'r', encoding='utf-8') as f:
                for line in f:
                    word = line.strip()
                    if word:
                        self.add_word(word, base_form=word, definition=get_def(word))
                        imported += 1
                        
        elif ext == '.csv':
            with open(file_path, 'r', encoding='utf-8') as f:
                reader = csv.reader(f)
                for row in reader:
                    if row:
                        word = row[0].strip()
                        definition = row[1].strip() if len(row) > 1 else ""
                        self.add_word(word, base_form=word, definition=get_def(word, definition))
                        imported += 1
                        
        elif ext == '.xlsx':
            from openpyxl import load_workbook
            wb = load_workbook(file_path)
            ws = wb.active
            for row in ws.iter_rows(values_only=True):
                if row and row[0]:
                    word = str(row[0]).strip()
                    definition = str(row[1]).strip() if len(row) > 1 and row[1] else ""
                    self.add_word(word, base_form=word, definition=get_def(word, definition))
                    imported += 1
        
        return imported

    def export_to_file(self, export_path, words_list, target_lang="zh-CN", format_type="xlsx"):
        if format_type == "xlsx":
            wb = Workbook()
            ws = wb.active
            ws.append(["单词", "释义", "掌握程度", "来源文章数", "添加时间"])
            
            for word_info in words_list:
                ws.append([
                    word_info["base_form"],
                    word_info["definition"],
                    ["未掌握", "熟悉", "已掌握"][word_info["mastery_level"]],
                    len(word_info["source_articles"]),
                    word_info["add_time"]
                ])
            wb.save(export_path)
            
        elif format_type == "csv":
            with open(export_path, 'w', encoding='utf-8-sig', newline='') as f:
                writer = csv.writer(f)
                writer.writerow(["单词", "释义", "掌握程度", "来源文章数"])
                for word_info in words_list:
                    writer.writerow([
                        word_info["base_form"],
                        word_info["definition"],
                        ["未掌握", "熟悉", "已掌握"][word_info["mastery_level"]],
                        len(word_info["source_articles"])
                    ])
                    
        elif format_type == "txt":
            with open(export_path, 'w', encoding='utf-8') as f:
                for word_info in words_list:
                    f.write(f"{word_info['base_form']}\t{word_info['definition']}\n")
